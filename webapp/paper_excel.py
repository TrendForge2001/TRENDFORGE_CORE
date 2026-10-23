"""Spreadsheet exports for canonical paper P&L and technical research results."""

from __future__ import annotations

import math
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pandas as pd


def _safe_spreadsheet_value(value):
    if isinstance(value, str) and (
        value.startswith(("\t", "\r", "\n"))
        or value.lstrip().startswith(("=", "+", "-", "@"))
    ):
        return "'" + value
    return value


def safe_spreadsheet_frame(frame: pd.DataFrame) -> pd.DataFrame:
    """Preserve numeric types while neutralizing formula text and headers."""
    result = frame.map(_safe_spreadsheet_value)
    result.columns = [_safe_spreadsheet_value(value) for value in frame.columns]
    return result


def safe_csv_bytes(frame: pd.DataFrame) -> bytes:
    return safe_spreadsheet_frame(frame).to_csv(index=False).encode("utf-8-sig")


def paper_statement_day(statement: dict) -> str:
    """Use the authoritative statement's IST date, including at midnight."""
    try:
        stamp = datetime.fromisoformat(
            str(statement["period_end"]).replace("Z", "+00:00")
        )
    except (KeyError, TypeError, ValueError):
        raise ValueError("Daily paper statement requires a valid period end") from None
    if stamp.tzinfo is None:
        stamp = stamp.replace(tzinfo=ZoneInfo("Asia/Kolkata"))
    return stamp.astimezone(ZoneInfo("Asia/Kolkata")).date().isoformat()


def validate_daily_statement(statement: dict, portfolio: dict) -> None:
    """Reject incomplete/malformed accounting instead of displaying invented zeroes."""
    if not isinstance(statement, dict) or not isinstance(portfolio, dict):
        raise ValueError("Invalid P&L/trade ledger response")
    if (
        statement.get("scope") != "paper_trading_closed_trades"
        or statement.get("period") != "daily"
    ):
        raise ValueError("Expected canonical daily paper-trading P&L statement")
    summary = statement.get("summary")
    if not isinstance(summary, dict):
        raise ValueError("Invalid daily P&L summary")
    for key in ("trade_count", "gross_profit", "gross_loss", "net_pnl"):
        value = summary.get(key)
        try:
            valid = not isinstance(value, bool) and math.isfinite(float(value))
        except (TypeError, ValueError, OverflowError):
            valid = False
        if not valid:
            raise ValueError("Invalid daily P&L summary")
    daily = statement.get("daily", [])
    closed = statement.get("trades", [])
    opened = portfolio.get("positions", [])
    if not all(
        isinstance(rows, list) and all(isinstance(row, dict) for row in rows)
        for rows in (daily, closed, opened)
    ):
        raise ValueError("Invalid P&L/trade ledger response")
    if type(summary["trade_count"]) is not int or summary["trade_count"] != len(closed):
        raise ValueError("Daily P&L trade count does not match its ledger")
    paper_statement_day(statement)


def paper_daily_xlsx(statement: dict, portfolio: dict) -> bytes:
    """Generate four sheets from the authoritative /pnl statement."""
    validate_daily_statement(statement, portfolio)
    summary = statement["summary"]
    daily = statement.get("daily", [])
    closed = statement.get("trades", [])
    opened = portfolio.get("positions", [])
    summary_rows = [
        ["Metric", "Value"],
        ["Statement period", "Daily (Asia/Kolkata)"],
        ["Statement date", paper_statement_day(statement)],
        ["Period end", str(statement["period_end"])],
        ["Closed trades", int(summary["trade_count"])],
        ["Realized profit (₹)", float(summary["gross_profit"])],
        ["Realized loss (₹)", float(summary["gross_loss"])],
        ["Realized net P&L (₹)", float(summary["net_pnl"])],
        ["Open positions", int(portfolio.get("open_positions") or 0)],
        ["Open unrealized P&L (₹)", float(portfolio.get("unrealized_pnl") or 0)],
        [
            "Accounting note",
            "Open unrealized P&L uses last saved marks, "
            "NOT intraday movement or proof of a fresh quote.",
        ],
        ["Date basis", str(statement.get("date_basis") or "")],
    ]
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        tables = {
            "Daily P&L": pd.DataFrame(summary_rows[1:], columns=summary_rows[0]),
            "Daily Closed Trades": pd.DataFrame(closed),
            "Open Paper Portfolio": pd.DataFrame(opened),
            "Daily Aggregates": pd.DataFrame(daily),
        }
        for name, frame in tables.items():
            if frame.empty and not len(frame.columns):
                frame = pd.DataFrame({"No records": []})
            frame = safe_spreadsheet_frame(frame)
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(fill_type="solid", fgColor="102E44")
            for index, column in enumerate(frame.columns, 1):
                sheet.column_dimensions[get_column_letter(index)].width = min(
                    35, max(16, len(str(column)) + 3)
                )
    return buf.getvalue()


def strategy_paper_xlsx(statement: dict, strategy_id: str | None = None) -> bytes:
    """Export every strategy or one strategy with its own closed/open ledger."""
    if statement.get("scope") != "strategy_paper_trading" or not isinstance(
        statement.get("rows"), list
    ):
        raise ValueError("Invalid strategy statement")
    rows = statement["rows"]
    if not rows or not all(isinstance(row, dict) for row in rows):
        raise ValueError("Invalid strategy statement")
    if strategy_id is not None:
        rows = [row for row in rows if row.get("strategy_id") == strategy_id]
        if not rows:
            raise ValueError("Unknown strategy")
    for row in rows:
        for key in ("realized_pnl", "unrealized_pnl", "total_pnl"):
            value = row.get(key)
            if type(value) not in {int, float} or not math.isfinite(value):
                raise ValueError("Invalid strategy P&L")
        for key in ("closed_trades", "open_positions"):
            if not isinstance(row.get(key), list) or not all(
                isinstance(item, dict) for item in row[key]
            ):
                raise ValueError("Invalid strategy ledger")
    summary = [
        {
            key: value
            for key, value in row.items()
            if key not in {"closed_trades", "open_positions"}
        }
        for row in rows
    ]
    closed = [
        {**trade, "strategy_id": row["strategy_id"]}
        for row in rows
        for trade in row["closed_trades"]
    ]
    opened = [
        {**trade, "strategy_id": row["strategy_id"]}
        for row in rows
        for trade in row["open_positions"]
    ]
    tables = {
        "Strategy P&L": summary,
        "Closed Trades": closed,
        "Open Positions": opened,
        "Statement Basis": [
            {
                key: statement.get(key)
                for key in ("period", "period_start", "period_end", "accounting")
            }
        ],
    }
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, records in tables.items():
            frame = (
                pd.DataFrame(records) if records else pd.DataFrame({"No records": []})
            )
            # Nested metadata is text in Excel, retaining spreadsheet escaping.
            frame = frame.map(
                lambda value: str(value) if isinstance(value, (dict, list)) else value
            )
            safe_spreadsheet_frame(frame).to_excel(writer, sheet_name=name, index=False)
            writer.sheets[name].freeze_panes = "A2"
    return buf.getvalue()
