"""Technical/paper downloads preserve accounting and never execute exported text."""

import csv
from copy import deepcopy
from io import BytesIO, StringIO

import pandas as pd
import pytest
from openpyxl import load_workbook

from webapp.paper_excel import (
    paper_daily_xlsx,
    paper_statement_day,
    safe_csv_bytes,
    validate_daily_statement,
)


def statement():
    return {
        "scope": "paper_trading_closed_trades",
        "period": "daily",
        "period_end": "2026-10-22T20:00:00Z",
        "summary": {
            "trade_count": 1,
            "gross_profit": 0,
            "gross_loss": 12.5,
            "net_pnl": -12.5,
        },
        "trades": [{"symbol": "TCS", "pnl": -12.5}],
        "daily": [{"date": "2026-10-23", "trade_count": 1, "net_pnl": -12.5}],
    }


def test_statement_date_uses_ist_instead_of_ui_calendar_or_utc():
    assert paper_statement_day(statement()) == "2026-10-23"
    data = paper_daily_xlsx(
        statement(), {"positions": [], "open_positions": 0, "unrealized_pnl": -1.25}
    )
    sheet = load_workbook(BytesIO(data))["Daily P&L"]
    values = {row[0]: row[1] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert values["Statement date"] == "2026-10-23"
    assert values["Realized net P&L (₹)"] == -12.5
    assert values["Open unrealized P&L (₹)"] == -1.25


@pytest.mark.parametrize(
    "change",
    [
        {"summary": []},
        {"summary": None},
        {"trades": ["TCS"]},
        {"daily": ["invalid"]},
        {"period": "weekly"},
        {"period_end": "invalid"},
        {
            "summary": {
                "trade_count": 2,
                "gross_profit": 0,
                "gross_loss": 0,
                "net_pnl": 0,
            }
        },
        {
            "summary": {
                "trade_count": "1.0",
                "gross_profit": 0,
                "gross_loss": 0,
                "net_pnl": 0,
            }
        },
        {
            "summary": {
                "trade_count": 1,
                "gross_profit": float("nan"),
                "gross_loss": 0,
                "net_pnl": 0,
            }
        },
    ],
)
def test_malformed_daily_reports_are_rejected_before_totals_or_download(change):
    report = {**statement(), **deepcopy(change)}
    with pytest.raises(ValueError):
        validate_daily_statement(report, {})
    with pytest.raises(ValueError):
        paper_daily_xlsx(report, {})


def test_exports_escape_formula_headers_and_values_but_preserve_numeric_losses():
    frame = pd.DataFrame(
        {
            '=HYPERLINK("https://example.invalid")': [" =1+1"],
            "pnl": [-12.5],
            "note": ["\n@SUM(1)"],
        }
    )
    csv_text = safe_csv_bytes(frame).decode("utf-8-sig")
    rows = list(csv.reader(StringIO(csv_text)))
    assert rows[0][0].startswith("'=HYPERLINK")
    assert "' =1+1" in csv_text
    assert "'\n@SUM(1)" in csv_text
    assert ",-12.5," in csv_text
    assert frame.iloc[0, 0] == " =1+1"  # Export leaves displayed/source data intact.

    report = statement()
    report["trades"] = frame.to_dict(orient="records")
    data = paper_daily_xlsx(report, {"positions": []})
    sheet = load_workbook(BytesIO(data))["Daily Closed Trades"]
    assert sheet["A1"].value.startswith("'=HYPERLINK")
    assert sheet["A1"].data_type == "s"
    assert sheet["A2"].value == "' =1+1"
    assert sheet["A2"].data_type == "s"
    assert sheet["B2"].value == -12.5
    assert sheet["B2"].data_type == "n"
