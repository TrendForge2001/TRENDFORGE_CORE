"""Paper automation admission, recovery and accounting without external calls."""

import asyncio
import secrets
import threading
from copy import deepcopy
from datetime import datetime, timedelta
from io import BytesIO
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pandas as pd
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.paper_automation_router import paper_automation_router
from providers.kite_paper_quotes import FreshKitePaperQuotes
from scanner.strategy_registry import list_strategies
from services.strategy_paper_automation import StrategyPaperAutomation
from services.strategy_paper_reporting import StrategyPaperReporting
from webapp.paper_excel import strategy_paper_xlsx

TZ = ZoneInfo("Asia/Kolkata")
NOW = datetime(2026, 10, 23, 10, 0, tzinfo=TZ)


class Repository:
    def __init__(self):
        self.saved = {}
        self.claims = set()
        self.fail = False

    def load_state(self, key):
        return {"payload": self.saved.get(key, {})}

    def save_state(self, key, kind, payload):
        if self.fail:
            raise RuntimeError("Do not expose credentials")
        self.saved[key] = deepcopy(payload)

    def claim_execution(self, **kwargs):
        key = (kwargs["scan_run_id"], kwargs["strategy_id"])
        if key in self.claims:
            return False
        self.claims.add(key)
        return True


def build(*, repo=None, now=NOW):
    repo = repo or Repository()
    health = {
        "enabled": True,
        "kill_switch": False,
        "can_open_new_positions": True,
        "entry_storage_policy": {
            "entry_allowed": True,
            "mode": "durable",
            "verified_persistent": True,
        },
    }
    positions = []
    paper = SimpleNamespace(
        _entry_lock=threading.RLock(),
        health=lambda: deepcopy(health),
        runtime=SimpleNamespace(snapshot=lambda: {"positions": positions}),
    )
    scans, executions, monitors, validations = [], [], [], []

    def scan(**kwargs):
        scans.append(kwargs)
        return {
            "strategy_ids": [
                row["strategy_id"] for row in list_strategies(mode=kwargs["mode"])
            ]
        }

    def execute(**kwargs):
        executions.append(kwargs)
        return {"status": "blocked", "blockers": ["validation_gate_not_passed"]}

    validation = {
        "status": "completed",
        "run_id": "validation-1",
        "request": {
            "universes": ["NIFTY 100"],
            "source": "all",
            "modes": ["intraday", "swing", "positional"],
        },
    }
    quote = FreshKitePaperQuotes(
        SimpleNamespace(
            quote=lambda instruments: {
                key: {"last_price": 100, "timestamp": now} for key in instruments
            }
        )
    )
    monitor = SimpleNamespace(tick=lambda **kwargs: monitors.append(kwargs))
    service = StrategyPaperAutomation(
        scanner=SimpleNamespace(scan=scan),
        bridge=SimpleNamespace(
            paper_auto_entry=SimpleNamespace(paper_service=paper, enabled=True),
            SUPPORTED_MODES={"intraday", "swing"},
            execute=execute,
        ),
        repository=repo,
        quote_provider=quote,
        paper_monitor=monitor,
        strategy_monitor=SimpleNamespace(tick=lambda **kwargs: monitors.append(kwargs)),
        validation_getter=lambda: validation,
        validation_starter=lambda: validations.append(True) or {"status": "running"},
        clock=lambda: now,
    )
    return SimpleNamespace(
        service=service,
        repo=repo,
        health=health,
        positions=positions,
        scans=scans,
        executions=executions,
        monitors=monitors,
        validation=validation,
        validations=validations,
        now=now,
    )


def test_default_off_and_restart_never_resumes_entries():
    env = build()
    assert env.service.tick()["status"] == "off"
    assert not env.scans
    env.service.configure(True)
    restarted = build(repo=env.repo)
    assert restarted.service.status()["enabled"] is False
    assert restarted.service.status()["resume_required"] is True
    assert restarted.service.tick()["status"] == "off"
    assert not restarted.scans


def test_all_saved_strategies_are_considered_and_slots_are_claimed_once():
    env = build()
    env.service.configure(True)
    cycle = env.service.tick()
    assert cycle["status"] == "watching_market"
    assert {row["strategy_id"] for row in cycle["strategies"]} == {
        row["strategy_id"] for row in list_strategies()
    }
    assert [row["mode"] for row in env.scans] == ["intraday", "swing", "positional"]
    assert all(
        row["universe"] == "NIFTY 100" and row["source"] == "all" for row in env.scans
    )
    assert any(
        row["blockers"] == ["paper_mode_not_supported"] for row in cycle["strategies"]
    )
    env.service.tick()
    assert len(env.scans) == 3
    assert all(
        callable(row["admission_guard"]) and callable(row["candidate_transform"])
        for row in env.executions
    )


@pytest.mark.parametrize(
    "field,value",
    [
        ("enabled", False),
        ("kill_switch", True),
        ("kill_switch", None),
        ("can_open_new_positions", False),
    ],
)
def test_runtime_blockers_precede_scans(field, value):
    env = build()
    env.health[field] = value
    env.service.configure(True)
    assert env.service.tick()["status"] == "blocked"
    assert not env.scans and not env.validations


@pytest.mark.parametrize(
    "policy",
    [
        {},
        {"entry_allowed": True, "mode": "ephemeral_simulation"},
        {"entry_allowed": True, "mode": "local_simulation"},
        {"entry_allowed": True, "mode": "durable", "verified_persistent": False},
    ],
)
def test_durable_storage_cannot_be_replaced_by_an_ephemeral_opt_in(policy):
    env = build()
    env.health["entry_storage_policy"] = policy
    env.service.configure(True)
    assert env.service.tick()["status"] == "blocked"
    assert not env.scans


@pytest.mark.parametrize(
    "change",
    [
        {"status": "not_started"},
        {"request": {"source": "template"}},
        {"request": {"universes": ["NIFTY 500"]}},
    ],
)
def test_validation_is_started_once_for_the_complete_catalog(change):
    env = build()
    env.validation.update(change)
    env.service.configure(True)
    assert env.service.tick()["status"] == "validating_strategies"
    env.service.tick()
    assert env.validations == [True]
    assert not env.scans


@pytest.mark.parametrize(
    "now", [NOW.replace(hour=8), NOW.replace(hour=15), NOW + timedelta(days=1)]
)
def test_no_entries_outside_market_window(now):
    env = build(now=now)
    env.service.configure(True)
    assert env.service.tick()["status"] == "waiting_for_market"
    assert not env.scans


def test_off_remains_effective_when_settings_write_fails():
    env = build()
    env.service.configure(True)
    env.repo.fail = True
    state = env.service.configure(False)
    assert state["enabled"] is False and state["settings_saved"] is False
    assert env.service.tick()["status"] == "off"
    with pytest.raises(RuntimeError, match="could not be saved"):
        env.service.configure(True)
    assert env.service.status()["enabled"] is False


def test_off_invalidates_inflight_admission_and_monitoring_continues():
    env = build()
    env.service.configure(True)
    generation = env.service._generation
    candidate = {
        "automation_quote_at": NOW.isoformat(),
        "entry_bar_at": NOW - timedelta(minutes=15),
        "technical_mode": "intraday",
    }
    assert env.service._admit(generation, candidate)
    env.service.configure(False)
    env.positions.append({"strategy_id": "tpl_supertrend_rsi"})
    assert not env.service._admit(generation, candidate)
    assert env.service.tick()["monitoring_existing"]
    assert len(env.monitors) == 2
    assert not env.scans
    env.service.configure(True)
    assert not env.service._admit(generation, candidate)


def test_live_candidate_uses_current_quote_and_rejects_stale_candles():
    env = build()
    candidate = {
        "symbol": "TCS",
        "technical_mode": "intraday",
        "entry_bar_at": NOW - timedelta(minutes=15),
        "signal": {"entry": 90, "stoploss": 95, "target2": 110},
        "engines": {"Risk Engine": {"metrics": {"entry": 90, "stoploss": 95}}},
    }
    current = env.service._live_candidate(deepcopy(candidate))
    assert current["signal"]["entry"] == 100
    assert current["engines"]["Risk Engine"]["metrics"]["entry"] == 100
    candidate["entry_bar_at"] = NOW - timedelta(minutes=45)
    with pytest.raises(ValueError, match="completed candle"):
        env.service._live_candidate(candidate)


def test_worker_lifecycle_is_bounded_and_default_off():
    async def check():
        env = build()
        assert await env.service.start()
        assert not await env.service.start()
        await asyncio.sleep(0.02)
        assert env.service.status()["worker_running"]
        assert await env.service.stop()
        assert not env.service.status()["worker_running"]
        assert not env.scans

    asyncio.run(check())


@pytest.mark.parametrize(
    "price,age",
    [
        (0, 0),
        (True, 0),
        (float("nan"), 0),
        (float("inf"), 0),
        (100, 61),
        (100, -6),
        (None, 0),
    ],
)
def test_invalid_or_stale_quotes_are_rejected(price, age):
    broker = SimpleNamespace(
        quote=lambda _: {
            "NSE:TCS": {"last_price": price, "timestamp": NOW - timedelta(seconds=age)}
        }
    )
    with pytest.raises(ValueError):
        FreshKitePaperQuotes(broker).prices(["NSE:TCS"], now=NOW)


@pytest.mark.parametrize(
    "stamp",
    [NOW, NOW.replace(tzinfo=None), NOW.astimezone(ZoneInfo("UTC")).isoformat()],
)
def test_current_quotes_accept_ist_naive_or_utc_timestamp(stamp):
    broker = SimpleNamespace(
        quote=lambda _: {"NSE:TCS": {"last_price": 100, "timestamp": stamp}}
    )
    assert (
        FreshKitePaperQuotes(broker).prices(["NSE:TCS"], now=NOW)["NSE:TCS"][
            "last_price"
        ]
        == 100
    )


def report_service(closed=(), positions=()):
    runtime = SimpleNamespace(
        _state_lock=threading.RLock(),
        snapshot=lambda: {"positions": list(positions)},
        trades=SimpleNamespace(closed_trades=lambda: list(closed)),
    )
    return StrategyPaperReporting(SimpleNamespace(runtime=runtime))


def test_strategy_statements_include_zero_trade_strategies_and_full_exit_date_ledger():
    key = list_strategies()[0]["strategy_id"]
    closed = [
        {
            "strategy_id": key,
            "created_at": "2026-10-20T09:30:00+05:30",
            "closed_at": "2026-10-23T04:00:00Z",
            "pnl": 2.5,
        }
        for _ in range(1005)
    ]
    closed += [
        {
            "strategy_id": key,
            "created_at": "2026-10-22T10:00:00+05:30",
            "closed_at": "2026-10-22T14:00:00+05:30",
            "pnl": 500,
        },
        {"strategy_id": key, "closed_at": "2026-10-24T10:00:00+05:30", "pnl": 500},
    ]
    result = report_service(closed, [{"strategy_id": key, "pnl": -25}]).statement(
        now=NOW
    )
    assert len(result["rows"]) == len(list_strategies())
    row = next(row for row in result["rows"] if row["strategy_id"] == key)
    assert row["trade_count"] == 1005 and row["realized_pnl"] == 2512.5
    assert row["unrealized_pnl"] == -25 and row["total_pnl"] == 2487.5
    assert all(
        row["trade_count"] == 0 for row in result["rows"] if row["strategy_id"] != key
    )


def test_orphan_and_manual_trades_are_not_dropped_and_exports_escape_formulas():
    trades = [
        {"strategy_id": key, "symbol": "=1+1", "closed_at": NOW.isoformat(), "pnl": -10}
        for key in ("archived", None)
    ]
    statement = report_service(trades).statement(now=NOW)
    assert {"archived", "unattributed"} <= {
        row["strategy_id"] for row in statement["rows"]
    }
    workbook = load_workbook(BytesIO(strategy_paper_xlsx(statement, "archived")))
    assert workbook.sheetnames == [
        "Strategy P&L",
        "Closed Trades",
        "Open Positions",
        "Statement Basis",
    ]
    assert all(
        cell.data_type != "f" for sheet in workbook for row in sheet for cell in row
    )
    assert any(
        cell.value == "'=1+1" for row in workbook["Closed Trades"] for cell in row
    )
    assert workbook["Strategy P&L"]["H2"].value is not None


@pytest.mark.parametrize("pnl", [float("nan"), float("inf"), True, None])
def test_bad_accounting_never_becomes_zero_or_an_export(pnl):
    with pytest.raises(ValueError):
        report_service([{"closed_at": NOW.isoformat(), "pnl": pnl}]).statement(now=NOW)


def test_api_requires_admin_and_strict_boolean_and_sanitizes_failures(monkeypatch):
    key = secrets.token_hex(24)
    monkeypatch.setenv("TRENDFORGE_ADMIN_API_KEY", key)
    env = build()
    app = FastAPI()
    app.include_router(
        paper_automation_router(
            automation_getter=lambda: env.service, paper_getter=lambda: None
        )
    )
    with TestClient(app) as client:
        assert client.get("/paper/automation").status_code == 401
        headers = {"X-TrendForge-Admin-Key": key}
        for bad in ("true", 1, None):
            assert (
                client.put(
                    "/paper/automation", json={"enabled": bad}, headers=headers
                ).status_code
                == 422
            )
        response = client.put(
            "/paper/automation", json={"enabled": True}, headers=headers
        )
        assert response.status_code == 200 and response.json()["enabled"] is True
        env.repo.fail = True
        failed = client.put(
            "/paper/automation", json={"enabled": True}, headers=headers
        )
        assert failed.status_code == 503 and "credentials" not in failed.text
        assert (
            client.get(
                "/paper/strategy-pnl?period=invalid", headers=headers
            ).status_code
            == 422
        )
        assert client.get("/paper/strategy-pnl", headers=headers).status_code == 503


def test_off_returns_while_quote_is_inflight_and_prevents_a_late_fill():
    from services.paper_auto_entry_service import PaperAutoEntryService

    env = build()
    env.service.paper.can_open_new_positions = True
    runtime = env.service.paper.runtime
    runtime.portfolio = SimpleNamespace(positions={})
    runtime.trades = SimpleNamespace(
        realized_pnl_for_entry_date=lambda _: 0,
        has_trade_on_entry_date=lambda *_: False,
    )
    runtime._can_open = lambda _: True
    opened = []
    env.service.paper.open = lambda *args, **kwargs: (
        opened.append(args) or {"trade_id": 1}
    )
    auto = PaperAutoEntryService(env.service.paper, enabled=True)
    waiting, release = threading.Event(), threading.Event()

    def transform(candidate):
        waiting.set()
        assert release.wait(timeout=5)
        candidate["automation_quote_at"] = NOW.isoformat()
        return candidate

    env.service.configure(True)
    generation = env.service._generation
    candidate = {
        "symbol": "TCS",
        "eligible": True,
        "entry_bar_at": NOW - timedelta(minutes=15),
        "technical_mode": "intraday",
        "signal": {
            "signal": "BUY",
            "overall_score": 100,
            "confidence": 100,
            "entry": 100,
            "stoploss": 95,
            "target2": 110,
        },
    }
    results = []
    worker = threading.Thread(
        target=lambda: results.append(
            auto.process_scan(
                {"outcome": "completed", "top_picks": [candidate]},
                run_id="race-1",
                execute_requested=True,
                now=NOW,
                candidate_transform=transform,
                admission_guard=lambda row: env.service._admit(generation, row),
            )
        )
    )
    worker.start()
    assert waiting.wait(timeout=5)
    try:
        assert env.service.configure(False)["enabled"] is False
    finally:
        release.set()
        worker.join(timeout=5)
    assert not worker.is_alive()
    assert not opened
    assert results[0]["opened_count"] == 0
    assert "paper_open_failed" in results[0]["evaluations"][0]["reasons"]


def test_paper_tab_controls_toggle_and_export_each_strategy(monkeypatch):
    from pathlib import Path

    from streamlit.testing.v1 import AppTest

    env = build()
    statement = report_service().statement(now=NOW)
    calls = []

    class Client:
        def get_json(self, path, **kwargs):
            if path == "/paper/automation":
                return env.service.status()
            if path == "/paper/strategy-pnl":
                return statement
            if path == "/workspace/paper-portfolio":
                return {
                    "portfolio": {
                        "open_positions": 0,
                        "positions": [],
                        "unrealized_pnl": 0,
                        "realized_pnl": 0,
                    }
                }
            if path == "/paper/readiness":
                return env.health
            raise AssertionError(path)

        def put_json(self, path, payload):
            calls.append((path, payload))
            return env.service.configure(payload["enabled"])

    # Existing daily reporting is independently unavailable in this focused test.
    client = Client()
    original = client.get_json

    def get_json(path, **kwargs):
        if path == "/pnl":
            from webapp.api_client import ApiError

            raise ApiError("Controlled daily ledger unavailable", status_code=503)
        return original(path, **kwargs)

    client.get_json = get_json
    monkeypatch.setattr("webapp.runtime.api_client", lambda **kwargs: client)
    monkeypatch.setattr("webapp.runtime.admin_key_configured", lambda: True)
    app = AppTest.from_file(
        Path(__file__).resolve().parents[1] / "webapp/pages/paper_session.py"
    ).run(timeout=15)
    assert not app.exception
    assert len(app.toggle) == 1 and not app.toggle[0].value
    app.toggle[0].set_value(True)
    next(
        item for item in app.button if item.label == "Apply automation switch"
    ).click().run(timeout=15)
    assert not app.exception
    assert calls[-1] == ("/paper/automation", {"enabled": True})
    assert len(app.get("download_button")) == 2
    next(
        item for item in app.button if item.label == "Stop automated paper entries"
    ).click().run(timeout=15)
    assert not app.exception
    assert calls[-1] == ("/paper/automation", {"enabled": False})
    assert not app.toggle[0].value


def test_multi_strategy_candles_drop_all_incomplete_and_future_rows():
    from services.multi_strategy_nifty_scanner import MultiStrategyNiftyScanner

    frame = pd.DataFrame(
        {"close": [100, 101, 102, 103]},
        index=pd.DatetimeIndex(
            [
                NOW - timedelta(minutes=15),
                NOW,
                NOW + timedelta(minutes=15),
                NOW + timedelta(minutes=30),
            ]
        ),
    )
    result = MultiStrategyNiftyScanner._completed_bars(frame, interval="15m", now=NOW)
    assert len(result) == 1 and result.index[0] == NOW - timedelta(minutes=15)
    with pytest.raises(ValueError, match="No completed candles"):
        MultiStrategyNiftyScanner._completed_bars(
            frame.iloc[1:], interval="15m", now=NOW
        )


def test_fill_rechecks_candle_freshness_after_waiting_for_admission():
    env = build()
    env.service.configure(True)
    candidate = {
        "entry_bar_at": NOW - timedelta(minutes=15),
        "technical_mode": "intraday",
        "automation_quote_at": (NOW + timedelta(minutes=14, seconds=59)).isoformat(),
    }
    env.service.clock = lambda: NOW + timedelta(minutes=15, seconds=1)
    assert not env.service._admit(env.service._generation, candidate)
